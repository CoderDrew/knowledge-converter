"""PPTX → document model.

The original application only inspected decks for reference images. This
reader converts each slide into a Section: a title, its text (paragraphs and
bulleted/numbered lists with nesting), tables, pictures, and speaker notes.

Shapes are read in visual order (top-to-bottom, then left-to-right), because
decks built from free text boxes have no meaningful XML order. When a slide has
no title placeholder, the most prominent short text near the top is used.
"""
from __future__ import annotations

from pathlib import Path
from typing import Any, Iterable

from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE, PP_PLACEHOLDER

from knowledge_engine.errors import ConversionError, InvalidDocumentError
from knowledge_engine.limits import Limits
from knowledge_engine.model import (
    Block, Document, DocumentStats, Heading, ImageRef, Inline, LineBreak, Link, ListBlock, ListItem, Notice, Paragraph,
    Section, Table, TableCell, TableRow, Text, has_visible_content, normalize_inlines, plain_text,
)
from knowledge_engine.readers.common import AssetRegistry, WarningLog, safe_href


A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
ROW_TOLERANCE_EMU = 228_600  # 0.25 inch: shapes this close vertically share a row
TITLE_BAND = 0.45  # untitled slides: look for a title in the top 45% of the slide
TITLE_TYPES = {PP_PLACEHOLDER.TITLE, PP_PLACEHOLDER.CENTER_TITLE, PP_PLACEHOLDER.VERTICAL_TITLE}
SKIPPED_PLACEHOLDERS = {PP_PLACEHOLDER.FOOTER, PP_PLACEHOLDER.DATE, PP_PLACEHOLDER.SLIDE_NUMBER, PP_PLACEHOLDER.HEADER}
BULLETED_PLACEHOLDERS = {PP_PLACEHOLDER.BODY, PP_PLACEHOLDER.OBJECT, PP_PLACEHOLDER.VERTICAL_BODY, PP_PLACEHOLDER.VERTICAL_OBJECT}
DIAGRAM_URI = "http://schemas.openxmlformats.org/drawingml/2006/diagram"


class PptxStageError(ConversionError):
    def __init__(self, slide: int, stage: str) -> None:
        self.slide = slide
        self.stage = stage
        super().__init__("conversion_failed", f"Conversion failed while processing slide {slide} ({stage}).")

    def diagnostic_context(self) -> dict[str, Any]:
        return {"stage": self.stage, "slide": self.slide}


def _placeholder_type(shape: Any) -> Any | None:
    if not getattr(shape, "is_placeholder", False):
        return None
    try:
        return shape.placeholder_format.type
    except (AttributeError, ValueError):
        return None


def _shape_type(shape: Any) -> Any | None:
    try:
        return shape.shape_type
    except (NotImplementedError, AttributeError, ValueError):
        return None


def _visual_order(shapes: Iterable[Any]) -> list[Any]:
    def key(item: tuple[int, Any]) -> tuple[int, int, int]:
        index, shape = item
        top = shape.top if shape.top is not None else 0
        left = shape.left if shape.left is not None else 0
        return (top // ROW_TOLERANCE_EMU, left, index)

    return [shape for _, shape in sorted(enumerate(shapes), key=key)]


class PptxReader:
    kind = "pptx"

    def __init__(self, source: Path, source_name: str, limits: Limits) -> None:
        self.source_name = source_name
        self.limits = limits
        try:
            self.prs = Presentation(str(source))
        except Exception:
            raise InvalidDocumentError("corrupted", "The presentation is damaged or uses structures that could not be opened.") from None
        self.warnings = WarningLog(limits.max_warnings)
        self.assets = AssetRegistry("pptx", limits, self.warnings)
        self.stats = DocumentStats()
        self.slide_number = 0
        self.stage = "opening PPTX"
        self._slide_image_count = 0
        self._slide_title: str | None = None
        self._words = 0

    def read(self) -> Document:
        sections: list[Block] = []
        for number, slide in enumerate(self.prs.slides, 1):
            self.slide_number = number
            try:
                sections.append(self._slide(slide))
            except ConversionError:
                raise
            except Exception as exc:
                raise PptxStageError(number, self.stage) from exc
        self.stats.slides = len(sections)
        self.stats.images = len(self.assets.assets)
        self.stats.words = self._words
        if not sections:
            self.warnings.add("The presentation does not contain any slides.")
        title = self._title(sections)
        # The deck title heads the document; each slide is a level-2 section.
        return Document(
            title=title,
            source_name=self.source_name,
            source_kind="pptx",
            blocks=[Heading(1, [Text(title)]), *sections],
            assets=self.assets.assets,
            warnings=self.warnings.finish(),
            stats=self.stats,
        )

    def _title(self, sections: list[Block]) -> str:
        core = (self.prs.core_properties.title or "").strip()
        if core:
            return core
        first = sections[0] if sections else None
        if isinstance(first, Section) and first.title:
            return plain_text(first.title)
        return Path(self.source_name).stem

    # -------------------------------------------------------------- slides

    def _slide(self, slide: Any) -> Section:
        self.stage = "slide title"
        self._slide_image_count = 0
        shapes = _visual_order(slide.shapes)
        title_shape = self._find_title(slide, shapes)
        title = None
        if title_shape is not None:
            title = normalize_inlines(self._title_inlines(title_shape))
            if not has_visible_content(title):
                title = None
        self._slide_title = plain_text(title) if title else None
        section = Section(number=self.slide_number, label=f"Slide {self.slide_number}", title=title)
        if title:
            self.stats.headings += 1
            self._words += len(plain_text(title).split())
        for shape in shapes:
            if shape is title_shape:
                continue
            section.blocks.extend(self._shape(shape))
        self.stage = "speaker notes"
        try:
            if slide.has_notes_slide:
                frame = slide.notes_slide.notes_text_frame
                if frame is not None:
                    for para in frame.paragraphs:
                        inlines = normalize_inlines(self._paragraph_inlines(para))
                        if has_visible_content(inlines):
                            section.notes.append(inlines)
        except Exception:
            self.warnings.add(f"Speaker notes on slide {self.slide_number} could not be read.")
        if getattr(slide, "_element", None) is not None and slide._element.get("show") == "0":
            self.warnings.add(f"Slide {self.slide_number} is hidden in the presentation; it is included in the output.")
        return section

    def _find_title(self, slide: Any, shapes: list[Any]) -> Any | None:
        for shape in shapes:
            if _placeholder_type(shape) in TITLE_TYPES and shape.has_text_frame and shape.text_frame.text.strip():
                return shape
        # Decks built from plain text boxes: pick the largest short,
        # single-paragraph text in the upper part of the slide (topmost wins ties).
        slide_height = self.prs.slide_height or 6_858_000
        best = None
        best_key = None
        for shape in shapes:
            if not getattr(shape, "has_text_frame", False) or _placeholder_type(shape) in SKIPPED_PLACEHOLDERS:
                continue
            paragraphs = [p for p in shape.text_frame.paragraphs if p.text.strip()]
            if len(paragraphs) != 1 or len(paragraphs[0].text.strip()) > 120:
                continue
            top = shape.top or 0
            if top > slide_height * TITLE_BAND:
                continue
            sizes = [run.font.size for run in paragraphs[0].runs if run.font.size]
            size = max(sizes) if sizes else 0
            key = (-size, top // ROW_TOLERANCE_EMU, shape.left or 0)
            if best_key is None or key < best_key:
                best, best_key = shape, key
        return best

    def _title_inlines(self, shape: Any) -> list[Inline]:
        inlines: list[Inline] = []
        for index, para in enumerate(p for p in shape.text_frame.paragraphs if p.text.strip()):
            if index:
                inlines.append(Text(" "))
            # Titles are headings; drop run emphasis so the heading stays clean.
            inlines.extend(Text(n.text) if isinstance(n, Text) else n for n in self._paragraph_inlines(para) if not isinstance(n, LineBreak))
        return inlines

    # -------------------------------------------------------------- shapes

    def _shape(self, shape: Any) -> list[Block]:
        self.stage = "shape"
        ph_type = _placeholder_type(shape)
        if ph_type in SKIPPED_PLACEHOLDERS:
            return []
        shape_type = _shape_type(shape)
        if shape_type == MSO_SHAPE_TYPE.GROUP:
            self.stage = "group shape"
            blocks: list[Block] = []
            for child in _visual_order(shape.shapes):
                blocks.extend(self._shape(child))
            return blocks
        if getattr(shape, "has_table", False) and shape.has_table:
            self.stage = "table"
            return [self._table(shape.table)]
        if getattr(shape, "has_chart", False) and shape.has_chart:
            self.warnings.add(f"A chart on slide {self.slide_number} cannot be converted and was omitted.")
            return [Paragraph([Notice("[Chart not included]")])]
        if shape._element.tag.endswith("}graphicFrame"):
            graphic = shape._element.find(f".//{A}graphicData")
            if graphic is not None and graphic.get("uri") == DIAGRAM_URI:
                self.warnings.add(f"A SmartArt diagram on slide {self.slide_number} cannot be converted and was omitted.")
                return [Paragraph([Notice("[SmartArt not included]")])]
        if shape_type in {MSO_SHAPE_TYPE.MEDIA, MSO_SHAPE_TYPE.EMBEDDED_OLE_OBJECT, MSO_SHAPE_TYPE.LINKED_OLE_OBJECT}:
            self.warnings.add(f"Embedded media or an object on slide {self.slide_number} was omitted.")
            return []
        if shape_type == MSO_SHAPE_TYPE.PICTURE or hasattr(shape, "image"):
            image = self._picture(shape)
            return [Paragraph([image])] if image else []
        if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
            self.stage = "text frame"
            return self._text_frame(shape, ph_type in BULLETED_PLACEHOLDERS)
        return []

    def _picture(self, shape: Any) -> Inline | None:
        self.stage = "picture"
        try:
            image = shape.image
            blob = image.blob
        except Exception:
            self.warnings.add(f"A linked or unreadable picture on slide {self.slide_number} was omitted.")
            return Notice("[Image not included]")
        alt = None
        try:
            alt = shape._element._nvXxPr.cNvPr.get("descr")
        except AttributeError:
            pass
        self._slide_image_count += 1
        asset = self.assets.add(
            blob,
            content_type=image.content_type,
            extension_hint=image.ext,
            stem=f"slide-{self.slide_number:03d}-image-{self._slide_image_count:02d}",
            position=self.slide_number,
            alt=alt,
            nearest_heading=self._slide_title,
            slide=self.slide_number,
        )
        return ImageRef(asset.id, (alt or "").strip())

    @staticmethod
    def _bullet(para: Any, default_bulleted: bool) -> str | None:
        ppr = para._p.pPr
        if ppr is not None:
            if ppr.find(f"{A}buNone") is not None:
                return None
            if ppr.find(f"{A}buAutoNum") is not None:
                return "number"
            if ppr.find(f"{A}buChar") is not None or ppr.find(f"{A}buBlip") is not None:
                return "bullet"
        return "bullet" if default_bulleted else None

    def _text_frame(self, shape: Any, default_bulleted: bool) -> list[Block]:
        blocks: list[Block] = []
        stack: list[tuple[int, ListBlock]] = []
        for para in shape.text_frame.paragraphs:
            inlines = normalize_inlines(self._paragraph_inlines(para))
            if not has_visible_content(inlines):
                continue
            self._words += len(plain_text(inlines).split())
            bullet = self._bullet(para, default_bulleted)
            if bullet is None:
                stack.clear()
                blocks.append(Paragraph(inlines))
                self.stats.paragraphs += 1
                continue
            ordered = bullet == "number"
            level = para.level or 0
            while stack and level < stack[-1][0]:
                stack.pop()
            if stack and level == stack[-1][0] and stack[-1][1].ordered != ordered:
                stack.pop()
            if not stack:
                root = ListBlock(ordered=ordered)
                blocks.append(root)
                stack.append((level, root))
                self.stats.lists += 1
            elif level > stack[-1][0]:
                parent = stack[-1][1]
                if not parent.items:
                    parent.items.append(ListItem([]))
                nested = ListBlock(ordered=ordered)
                parent.items[-1].children.append(nested)
                stack.append((level, nested))
                self.stats.lists += 1
            stack[-1][1].items.append(ListItem(inlines))
            self.stats.list_items += 1
        return blocks

    def _paragraph_inlines(self, para: Any) -> list[Inline]:
        chunks: list[Inline] = []
        for child in para._p:
            tag = child.tag
            if tag in {f"{A}r", f"{A}fld"}:
                text_el = child.find(f"{A}t")
                text = text_el.text if text_el is not None and text_el.text else ""
                rpr = child.find(f"{A}rPr")
                fmt = self._run_format(rpr)
                node: Inline = Text(text, **fmt)
                if rpr is not None:
                    click = rpr.find(f"{A}hlinkClick")
                    if click is not None:
                        rid = click.get("{http://schemas.openxmlformats.org/officeDocument/2006/relationships}id")
                        href = None
                        if rid:
                            try:
                                href = safe_href(para.part.rels[rid].target_ref)
                            except (KeyError, AttributeError):
                                href = None
                        if href:
                            node = Link(href, [node])
                chunks.append(node)
            elif tag == f"{A}br":
                chunks.append(LineBreak())
        return chunks

    @staticmethod
    def _run_format(rpr: Any) -> dict[str, Any]:
        fmt: dict[str, Any] = {"bold": False, "italic": False, "underline": False, "strike": False, "vertical": None}
        if rpr is None:
            return fmt
        fmt["bold"] = rpr.get("b") in {"1", "true"}
        fmt["italic"] = rpr.get("i") in {"1", "true"}
        fmt["underline"] = rpr.get("u") not in {None, "none"}
        fmt["strike"] = rpr.get("strike") not in {None, "noStrike"}
        baseline = rpr.get("baseline")
        if baseline and baseline.lstrip("-").isdigit() and int(baseline) != 0:
            fmt["vertical"] = "sup" if int(baseline) > 0 else "sub"
        return fmt

    # -------------------------------------------------------------- tables

    def _table(self, table: Any) -> Table:
        rows: list[TableRow] = []
        has_header = bool(getattr(table, "first_row", True))
        for row_index, row in enumerate(table.rows):
            cells: list[TableCell] = []
            for cell in row.cells:
                if cell.is_spanned:
                    continue
                lines = []
                for para in cell.text_frame.paragraphs:
                    inlines = normalize_inlines(self._paragraph_inlines(para))
                    if has_visible_content(inlines):
                        lines.append(inlines)
                        self._words += len(plain_text(inlines).split())
                span_w = cell.span_width if cell.is_merge_origin else 1
                span_h = cell.span_height if cell.is_merge_origin else 1
                cells.append(TableCell(lines, colspan=span_w, rowspan=span_h))
            rows.append(TableRow(cells, header=has_header and row_index == 0))
        self.stats.tables += 1
        return Table(rows)
