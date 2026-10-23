"""DOCX → document model.

Ported from the FS Knowledge Converter's DocxConverter. The body-order
traversal, list-numbering resolution (including style-level numbering and
graceful recovery from incomplete numbering definitions), ordered-list
continuation across image-only paragraphs, callout styles, and the
nearest-heading/step image context all carry over. Additions: nested lists,
merged table cells, content controls, tracked changes, AlternateContent,
legacy VML images, underline/strike/super/subscript, and authored alt text.
"""
from __future__ import annotations

import re
from pathlib import Path
from typing import Any, Iterable

from docx import Document as open_docx
from docx.oxml.ns import qn
from docx.text.paragraph import Paragraph as DocxParagraph

from knowledge_engine.errors import ConversionError, InvalidDocumentError
from knowledge_engine.limits import Limits
from knowledge_engine.model import (
    Document, DocumentStats, Heading, ImageRef, Inline, LineBreak, Link, ListBlock, ListItem, Notice, Paragraph,
    Table, TableCell, TableRow, Text, has_visible_content, normalize_inlines, plain_text,
)
from knowledge_engine.readers.common import AssetRegistry, WarningLog, safe_href


A_NS = "http://schemas.openxmlformats.org/drawingml/2006/main"
WP_NS = "http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing"
MC_NS = "http://schemas.openxmlformats.org/markup-compatibility/2006"
V_NS = "urn:schemas-microsoft-com:vml"
R_NS = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS = {"a": A_NS, "wp": WP_NS, "v": V_NS, "w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
MC_ALTERNATE = f"{{{MC_NS}}}AlternateContent"
MC_CHOICE = f"{{{MC_NS}}}Choice"
MC_FALLBACK = f"{{{MC_NS}}}Fallback"
V_IMAGEDATA = f"{{{V_NS}}}imagedata"
R_ID = f"{{{R_NS}}}id"
O_TITLE = "{urn:schemas-microsoft-com:office:office}title"

# Containers whose children belong to the surrounding flow.
BLOCK_CONTAINERS = {qn("w:customXml")}
INLINE_CONTAINERS = {qn("w:ins"), qn("w:smartTag"), qn("w:customXml"), qn("w:fldSimple"), qn("w:moveTo"), qn("w:bdo"), qn("w:dir")}
SKIPPED_INLINE = {qn("w:del"), qn("w:moveFrom")}
CALLOUT_TERMS = ("warning", "caution", "tip", "note")


class DocxStageError(ConversionError):
    """Unexpected failure enriched with non-content location context."""

    def __init__(self, *, stage: str, element_index: int, element_type: str, paragraph_index: int | None, style: str | None) -> None:
        self.stage = stage
        self.element_index = element_index
        self.element_type = element_type
        self.paragraph_index = paragraph_index
        self.paragraph_style = style
        where = f"paragraph {paragraph_index}" if paragraph_index is not None else f"document element {element_index}"
        super().__init__("conversion_failed", f"Conversion failed while processing {where} ({stage}).")

    def diagnostic_context(self) -> dict[str, Any]:
        return {
            "stage": self.stage,
            "document_element_index": self.element_index,
            "element_type": self.element_type,
            "paragraph_index": self.paragraph_index,
            "paragraph_style": self.paragraph_style,
        }


def _flag(rpr: Any, tag: str) -> bool:
    element = rpr.find(qn(tag))
    if element is None:
        return False
    value = element.get(qn("w:val"))
    return value is None or value.lower() not in {"0", "false", "off", "no", "none"}


class DocxReader:
    kind = "docx"

    def __init__(self, source: Path, source_name: str, limits: Limits) -> None:
        self.source = source
        self.source_name = source_name
        self.limits = limits
        try:
            self.doc = open_docx(str(source))
        except Exception:
            raise InvalidDocumentError("corrupted", "The Word document is damaged or uses structures that could not be opened.") from None
        self.warnings = WarningLog(limits.max_warnings)
        self.assets = AssetRegistry("docx", limits, self.warnings)
        self.stats = DocumentStats()
        self.blocks: list[Any] = []
        self.position = 0
        self.paragraph_index = 0
        self.table_index = 0
        self.stage = "opening DOCX"
        self.nearest_heading: str | None = None
        self.nearest_step: str | None = None
        self._words = 0
        # list state
        self._stack: list[tuple[int, ListBlock]] = []
        self._ordered_count = 0
        self._ordered_can_continue = False
        self._rel_assets: dict[str, str] = {}

    # ------------------------------------------------------------- traversal

    def _body_elements(self, parent: Any) -> Iterable[Any]:
        for child in parent.iterchildren():
            if child.tag in (qn("w:p"), qn("w:tbl")):
                yield child
            elif child.tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    yield from self._body_elements(content)
            elif child.tag in BLOCK_CONTAINERS:
                yield from self._body_elements(child)

    def read(self) -> Document:
        for element in self._body_elements(self.doc.element.body):
            self.position += 1
            is_paragraph = element.tag == qn("w:p")
            paragraph = DocxParagraph(element, self.doc) if is_paragraph else None
            if is_paragraph:
                self.paragraph_index += 1
                self.stage = "paragraph initialization"
            else:
                self.table_index += 1
                self.stage = "table initialization"
            try:
                if paragraph is not None:
                    self._convert_paragraph(paragraph)
                else:
                    self._close_list()
                    self._convert_table(element)
            except ConversionError:
                raise
            except Exception as exc:
                raise DocxStageError(
                    stage=self.stage,
                    element_index=self.position,
                    element_type="paragraph" if is_paragraph else "table",
                    paragraph_index=self.paragraph_index if is_paragraph else None,
                    style=self._style_name(paragraph) if paragraph is not None else None,
                ) from exc
        self._close_list()
        if not self.blocks:
            self.warnings.add("The document did not contain convertible body content.")
        self.stats.images = len(self.assets.assets)
        self.stats.words = self._words
        return Document(
            title=self._title(),
            source_name=self.source_name,
            source_kind="docx",
            blocks=self.blocks,
            assets=self.assets.assets,
            warnings=self.warnings.finish(),
            stats=self.stats,
        )

    def _title(self) -> str:
        for block in self.blocks:
            if isinstance(block, Heading) and block.level == 1:
                text = plain_text(block.inlines)
                if text:
                    return text
        core_title = (self.doc.core_properties.title or "").strip()
        return core_title or Path(self.source_name).stem

    # ---------------------------------------------------------- paragraphs

    @staticmethod
    def _style_name(paragraph: DocxParagraph | None) -> str | None:
        if paragraph is None:
            return None
        try:
            style = paragraph.style
            name = style.name if style is not None else None
            return str(name) if name else None
        except (AttributeError, KeyError, ValueError):
            return None

    def _warn_numbering(self, reason: str) -> None:
        self.warnings.add(
            f"Could not fully resolve list formatting at document position {self.position}: {reason}. "
            "Content was preserved using the paragraph style when possible."
        )

    def _style_num_pr(self, paragraph: DocxParagraph) -> Any | None:
        try:
            style = paragraph.style
            style_element = getattr(style, "element", None) if style is not None else None
            return getattr(getattr(style_element, "pPr", None), "numPr", None)
        except (AttributeError, KeyError, ValueError):
            self._warn_numbering("the paragraph style numbering definition was unavailable")
            return None

    def _paragraph_kind(self, paragraph: DocxParagraph) -> tuple[str, int | None]:
        self.stage = "paragraph style"
        style = (self._style_name(paragraph) or "").lower()
        match = re.match(r"heading\s*(\d+)", style)
        if match:
            return "heading", max(1, min(int(match.group(1)), 6))
        if style == "title":
            return "heading", 1
        # Word often stores numbering on the paragraph style rather than on the
        # paragraph itself (the built-in List Number/List Bullet styles do this).
        self.stage = "list numbering"
        num_pr = getattr(getattr(paragraph._p, "pPr", None), "numPr", None)
        if num_pr is None:
            num_pr = self._style_num_pr(paragraph)
        if num_pr is not None:
            num_id_element = getattr(num_pr, "numId", None)
            num_id_value = num_id_element.get(qn("w:val")) if num_id_element is not None else None
            level_element = getattr(num_pr, "ilvl", None)
            level_value = level_element.get(qn("w:val")) if level_element is not None else None
            if num_id_value is None:
                self._warn_numbering("w:numPr has no usable w:numId value")
                return self._style_list_fallback(style)
            level = 0
            try:
                num_id = int(num_id_value)
                level = int(level_value) if level_value is not None else 0
                if level_value is None and level_element is not None:
                    self._warn_numbering("w:ilvl has no value; level 0 was used")
                if num_id == 0:
                    # numId 0 explicitly removes numbering inherited from the style.
                    return "paragraph", None
                numbering = getattr(getattr(self.doc.part, "numbering_part", None), "element", None)
                if numbering is None:
                    raise LookupError("numbering part is absent")
                num_el = numbering.xpath(f'./w:num[@w:numId="{num_id}"]')[0]
                abstract_id = num_el.xpath("./w:abstractNumId/@w:val")[0]
                fmt = numbering.xpath(
                    f'./w:abstractNum[@w:abstractNumId="{abstract_id}"]/w:lvl[@w:ilvl="{level}"]/w:numFmt/@w:val'
                )[0]
                return ("bullet" if fmt in {"bullet", "none"} else "number"), level
            except (IndexError, AttributeError, KeyError, LookupError, TypeError, ValueError) as exc:
                self._warn_numbering(f"numbering definition is incomplete ({type(exc).__name__})")
                return self._style_list_fallback(style, level)
        # Some producers omit numPr but preserve a recognizable list style.
        if "list bullet" in style:
            return "bullet", 0
        if "list number" in style:
            return "number", 0
        return "paragraph", None

    @staticmethod
    def _style_list_fallback(style: str, level: int = 0) -> tuple[str, int | None]:
        if "bullet" in style:
            return "bullet", level
        if "number" in style:
            return "number", level
        return "paragraph", None

    def _count_words(self, text: str) -> None:
        self._words += len(text.split())

    def _convert_paragraph(self, paragraph: DocxParagraph) -> None:
        kind, level = self._paragraph_kind(paragraph)
        self.stage = "paragraph text"
        text = " ".join(paragraph.text.split())
        if kind == "heading":
            self._close_list()
            self._ordered_can_continue = False
            if text:
                self.nearest_heading = text
                self.stage = "heading inline content"
                inlines = normalize_inlines(self._inline_content(paragraph._p))
                self.blocks.append(Heading(level or 1, inlines))
                self.stats.headings += 1
                self._count_words(text)
            return
        if kind in {"number", "bullet"}:
            if text:
                self.nearest_step = self._step_label(text, kind)
            self.stage = "list item inline content"
            inlines = normalize_inlines(self._inline_content(paragraph._p))
            self._add_list_item(kind == "number", level or 0, inlines)
            self._count_words(text)
            return
        self._close_list()
        self.stage = "paragraph inline content"
        inlines = normalize_inlines(self._inline_content(paragraph._p))
        if not has_visible_content(inlines):
            return
        style = (self._style_name(paragraph) or "").lower()
        callout = next((term for term in CALLOUT_TERMS if term in style), None)
        self.blocks.append(Paragraph(inlines, callout))
        self.stats.paragraphs += 1
        self._count_words(text)
        # An image-only paragraph commonly separates two numbered Word steps.
        # Preserve the next visible ordinal when that happens.
        if text or not any(isinstance(node, ImageRef) for node in inlines):
            self._ordered_can_continue = False

    def _step_label(self, text: str, kind: str) -> str | None:
        match = re.match(r"(?i)(step\s+\d+[:.]?)", text)
        if match:
            return match.group(1).rstrip(".:")
        if kind == "number":
            return text[:100]
        return self.nearest_step

    # --------------------------------------------------------------- lists

    def _new_list(self, ordered: bool) -> ListBlock:
        self.stats.lists += 1
        return ListBlock(ordered=ordered)

    def _add_list_item(self, ordered: bool, level: int, inlines: list[Inline]) -> None:
        while self._stack and level < self._stack[-1][0]:
            self._stack.pop()
        if self._stack and level == self._stack[-1][0] and self._stack[-1][1].ordered != ordered:
            if len(self._stack) == 1:
                self._close_list()
            else:
                self._stack.pop()
        if not self._stack:
            root = self._new_list(ordered)
            if ordered:
                if self._ordered_can_continue and self._ordered_count:
                    root.start = self._ordered_count + 1
                else:
                    self._ordered_count = 0
            self.blocks.append(root)
            self._stack.append((level, root))
        elif level > self._stack[-1][0]:
            parent = self._stack[-1][1]
            if not parent.items:
                parent.items.append(ListItem([]))
            nested = self._new_list(ordered)
            parent.items[-1].children.append(nested)
            self._stack.append((level, nested))
        current = self._stack[-1][1]
        current.items.append(ListItem(inlines))
        self.stats.list_items += 1
        if len(self._stack) == 1:
            if current.ordered:
                self._ordered_count += 1
                self._ordered_can_continue = True
            else:
                self._ordered_can_continue = False

    def _close_list(self) -> None:
        self._stack.clear()

    # ------------------------------------------------------------- inlines

    def _inline_content(self, container: Any) -> list[Inline]:
        chunks: list[Inline] = []
        for child in container:
            tag = child.tag
            if tag == qn("w:r"):
                chunks.extend(self._run_content(child))
            elif tag == qn("w:hyperlink"):
                self.stage = "hyperlink relationship"
                label = self._inline_content(child)
                href = self._hyperlink_target(child)
                if href:
                    chunks.append(Link(href, label))
                else:
                    chunks.extend(label)
            elif tag == qn("w:sdt"):
                content = child.find(qn("w:sdtContent"))
                if content is not None:
                    chunks.extend(self._inline_content(content))
            elif tag in INLINE_CONTAINERS:
                chunks.extend(self._inline_content(child))
            elif tag == MC_ALTERNATE:
                chunks.extend(self._alternate_content(child, self._inline_content))
            elif tag in SKIPPED_INLINE:
                continue
        return chunks

    def _hyperlink_target(self, element: Any) -> str | None:
        rid = element.get(qn("r:id"))
        anchor = element.get(qn("w:anchor"))
        if rid:
            rel = self.doc.part.rels.get(rid)
            target = getattr(rel, "target_ref", None) if rel is not None else None
            if not target:
                self.warnings.add(f"A hyperlink at document position {self.position} has no target; its label was preserved.")
                return None
            href = safe_href(str(target) + (f"#{anchor}" if anchor else ""))
            if href is None:
                self.warnings.add(f"A hyperlink at document position {self.position} uses an unsupported scheme and was removed.")
            return href
        # Internal bookmark targets do not exist in the HTML; keep the label only.
        return None

    def _alternate_content(self, element: Any, handler: Any) -> list[Inline]:
        for branch in (element.find(MC_CHOICE), element.find(MC_FALLBACK)):
            if branch is not None:
                result = handler(branch)
                if result:
                    return result
        return []

    def _run_content(self, run_el: Any, fmt: dict[str, Any] | None = None) -> list[Inline]:
        fmt = fmt or self._run_format(run_el)
        chunks: list[Inline] = []
        for child in run_el:
            tag = child.tag
            if tag == qn("w:t"):
                chunks.append(Text(child.text or "", **fmt))
            elif tag == qn("w:tab"):
                chunks.append(Text(" ", **fmt))
            elif tag == qn("w:noBreakHyphen"):
                chunks.append(Text("-", **fmt))
            elif tag == qn("w:br"):
                if child.get(qn("w:type")) not in {"page", "column"}:
                    chunks.append(LineBreak())
            elif tag == qn("w:cr"):
                chunks.append(LineBreak())
            elif tag == qn("w:drawing"):
                self.stage = "drawing relationship"
                chunks.extend(self._drawing(child))
            elif tag in {qn("w:pict"), qn("w:object")}:
                chunks.extend(self._vml(child))
            elif tag == MC_ALTERNATE:
                chunks.extend(self._alternate_content(child, lambda branch: self._run_content(branch, fmt)))
            elif tag in {qn("w:footnoteReference"), qn("w:endnoteReference")}:
                self.warnings.add("Footnotes and endnotes are not included in the output.")
            elif tag == qn("w:commentReference"):
                continue
        return chunks

    def _run_format(self, run_el: Any) -> dict[str, Any]:
        rpr = run_el.find(qn("w:rPr"))
        fmt: dict[str, Any] = {"bold": False, "italic": False, "underline": False, "strike": False, "vertical": None}
        if rpr is None:
            return fmt
        fmt["bold"] = _flag(rpr, "w:b")
        fmt["italic"] = _flag(rpr, "w:i")
        fmt["underline"] = _flag(rpr, "w:u")
        fmt["strike"] = _flag(rpr, "w:strike") or _flag(rpr, "w:dstrike")
        vert = rpr.find(qn("w:vertAlign"))
        if vert is not None:
            value = vert.get(qn("w:val"))
            fmt["vertical"] = "sup" if value == "superscript" else "sub" if value == "subscript" else None
        style = rpr.find(qn("w:rStyle"))
        style_id = (style.get(qn("w:val")) or "").lower() if style is not None else ""
        if "strong" in style_id:
            fmt["bold"] = True
        if "emphasis" in style_id:
            fmt["italic"] = True
        return fmt

    def _drawing(self, drawing: Any) -> list[Inline]:
        doc_pr = drawing.find(".//wp:docPr", NS)
        alt = None
        if doc_pr is not None:
            alt = doc_pr.get("descr") or doc_pr.get("title")
        blips = drawing.findall(".//a:blip", NS)
        chunks: list[Inline] = []
        for blip in blips:
            rid = blip.get(qn("r:embed"))
            if rid:
                chunks.append(self._image(rid, alt))
            elif blip.get(qn("r:link")):
                self.warnings.add(f"A linked (not embedded) image at document position {self.position} cannot be included.")
                chunks.append(Notice("[Linked image not included]"))
        if not blips:
            if drawing.find(".//w:txbxContent", NS) is not None:
                self.warnings.add(f"Text box content at document position {self.position} is not included.")
            else:
                self.warnings.add(f"A drawing, chart, or SmartArt graphic at document position {self.position} could not be represented.")
                chunks.append(Notice("[Unsupported drawing]"))
        return chunks

    def _vml(self, element: Any) -> list[Inline]:
        chunks: list[Inline] = []
        for image_data in element.iter(V_IMAGEDATA):
            rid = image_data.get(R_ID)
            if rid:
                chunks.append(self._image(rid, image_data.get(O_TITLE)))
        if not chunks:
            self.warnings.add(f"A legacy VML graphic at document position {self.position} may not be represented.")
        elif element.tag == qn("w:object"):
            self.warnings.add(f"An embedded object at document position {self.position} is shown as its preview image.")
        return chunks

    def _image(self, rid: str, alt: str | None) -> Inline:
        self.stage = "image relationship"
        try:
            part = self.doc.part.related_parts[rid]
            blob = getattr(part, "blob", None)
        except (KeyError, AttributeError):
            self.warnings.add(f"Missing image relationship {rid} at document position {self.position}.")
            return Notice("[Missing image]")
        if not isinstance(blob, (bytes, bytearray)):
            self.warnings.add(f"Image relationship {rid} has no readable image data at document position {self.position}.")
            return Notice("[Unreadable image]")
        asset = self.assets.add(
            bytes(blob),
            content_type=getattr(part, "content_type", None),
            extension_hint=Path(str(getattr(part, "partname", "") or "")).suffix,
            stem=f"image-{len(self.assets.assets) + 1:03d}",
            position=self.position,
            alt=alt,
            nearest_heading=self.nearest_heading,
            nearest_step=self.nearest_step,
            relationship_id=rid,
        )
        return ImageRef(asset.id, (alt or "").strip())

    # -------------------------------------------------------------- tables

    def _cell_lines(self, tc: Any) -> list[list[Inline]]:
        lines: list[list[Inline]] = []
        for element in self._body_elements(tc):
            if element.tag == qn("w:p"):
                self.stage = f"table {self.table_index} cell inline content"
                inlines = normalize_inlines(self._inline_content(element))
                if has_visible_content(inlines):
                    lines.append(inlines)
                    self._count_words(plain_text(inlines))
            else:
                self.warnings.add(f"A nested table at document position {self.position} was flattened into its cell.")
                for row in element.findall(qn("w:tr")):
                    for inner in row.findall(qn("w:tc")):
                        lines.extend(self._cell_lines(inner))
        return lines

    def _convert_table(self, tbl: Any) -> None:
        self.stage = "table structure"
        self._ordered_can_continue = False
        rows: list[TableRow] = []
        # grid column -> origin cell, for vertical merges
        open_merges: dict[int, TableCell] = {}
        explicit_headers = False
        for row_index, tr in enumerate(tbl.findall(qn("w:tr"))):
            self.stage = f"table {self.table_index} row {row_index + 1}"
            trpr = tr.find(qn("w:trPr"))
            is_header = trpr is not None and _flag(trpr, "w:tblHeader")
            explicit_headers = explicit_headers or is_header
            column = 0
            if trpr is not None:
                before = trpr.find(qn("w:gridBefore"))
                if before is not None and (before.get(qn("w:val")) or "").isdigit():
                    column = int(before.get(qn("w:val")))
            cells: list[TableCell] = []
            for tc in tr.findall(qn("w:tc")):
                tcpr = tc.find(qn("w:tcPr"))
                span = 1
                merge = None
                if tcpr is not None:
                    grid_span = tcpr.find(qn("w:gridSpan"))
                    if grid_span is not None and (grid_span.get(qn("w:val")) or "").isdigit():
                        span = max(1, int(grid_span.get(qn("w:val"))))
                    v_merge = tcpr.find(qn("w:vMerge"))
                    if v_merge is not None:
                        merge = v_merge.get(qn("w:val")) or "continue"
                if merge == "continue" and column in open_merges:
                    open_merges[column].rowspan += 1
                else:
                    cell = TableCell(self._cell_lines(tc), colspan=span)
                    cells.append(cell)
                    if merge == "restart":
                        open_merges[column] = cell
                    else:
                        open_merges.pop(column, None)
                column += span
            rows.append(TableRow(cells, header=is_header))
        if not rows:
            return
        if not explicit_headers:
            # Word does not consistently record author intent; like the
            # original converter, treat the first row as the header.
            rows[0].header = True
        self.blocks.append(Table(rows))
        self.stats.tables += 1
