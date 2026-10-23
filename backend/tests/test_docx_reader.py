from pathlib import Path

import pytest
from docx import Document
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image

from knowledge_engine.model import Heading, ListBlock, Table
from knowledge_engine.pipeline import read_document
from knowledge_engine.profiles import PROFILES


def fresh_html(path: Path) -> str:
    return PROFILES["freshservice"].build(read_document(path, path.name)).html


# ---- carried over from the original converter's test suite ---------------


def test_docx_semantics_images_and_manifest(sample_docx: Path):
    document = read_document(sample_docx, sample_docx.name)
    article = PROFILES["freshservice"].build(document).html
    assert "<h1>Cash Management</h1>" in article
    assert "<strong>Cash Management</strong>" in article
    assert "<em> to review </em>" in article
    assert "<ol>" in article and "Step 1" in article
    assert "<ul>" in article and "Confirm the displayed balance" in article
    assert "<table>" in article and "<th>Field</th>" in article
    assert 'class="callout tip"' in article
    assert 'src="images/image-001.png"' in article
    assert article.index("Step 1") < article.index("image-001.png") < article.index("Step 2")
    asset = document.assets[0]
    assert asset.position == 3
    assert asset.nearest_heading == "Cash Management"
    assert asset.nearest_step.startswith("Step 1")
    assert len(asset.sha256) == 64
    assert document.stats.headings == 1
    assert document.stats.tables == 1
    assert document.stats.lists == 2
    assert document.title == "Cash Management"


def test_numbering_continues_across_image_only_paragraph(tmp_path: Path):
    png = tmp_path / "between.png"
    Image.new("RGB", (320, 180), "white").save(png)
    doc = Document()
    doc.add_paragraph("Step 1: Do this.", style="List Number")
    doc.add_paragraph().add_run().add_picture(str(png), width=Inches(2))
    doc.add_paragraph("Step 2: Do that.", style="List Number")
    path = tmp_path / "ordered.docx"; doc.save(path)
    assert '<ol start="2">' in fresh_html(path)


def _malformed_numbering_docx(tmp_path: Path, case: str) -> Path:
    doc = Document()
    paragraph = doc.add_paragraph("Preserve this procedural instruction.", style="List Number")
    ppr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    if case != "missing-num-id":
        num_id = OxmlElement("w:numId")
        if case != "empty-num-id":
            style_num_id = paragraph.style.element.pPr.numPr.numId.val
            num_id.set(qn("w:val"), "99999" if case == "missing-definition" else str(style_num_id))
        num_pr.append(num_id)
    if case == "empty-level":
        num_pr.append(OxmlElement("w:ilvl"))
    ppr.append(num_pr)
    path = tmp_path / f"{case}.docx"
    doc.save(path)
    return path


@pytest.mark.parametrize("case", ["missing-num-id", "empty-num-id", "empty-level", "missing-definition"])
def test_incomplete_numbering_properties_are_preserved_with_warning(tmp_path: Path, case: str):
    document = read_document(_malformed_numbering_docx(tmp_path, case), f"{case}.docx")
    article = PROFILES["freshservice"].build(document).html
    assert "Preserve this procedural instruction." in article
    assert "<ol" in article
    assert any("list formatting" in warning for warning in document.warnings)


def test_explicit_false_run_properties_are_not_rendered(tmp_path: Path):
    doc = Document()
    run = doc.add_paragraph().add_run("Not bold or italic")
    rpr = run._r.get_or_add_rPr()
    bold = OxmlElement("w:b"); bold.set(qn("w:val"), "0"); rpr.append(bold)
    italic = OxmlElement("w:i"); italic.set(qn("w:val"), "false"); rpr.append(italic)
    path = tmp_path / "false-properties.docx"; doc.save(path)
    article = fresh_html(path)
    assert "<strong>" not in article
    assert "<em>" not in article


# ---- improvements over the original -------------------------------------


def test_nested_lists_are_nested(rich_docx: Path):
    document = read_document(rich_docx, rich_docx.name)
    root = next(b for b in document.blocks if isinstance(b, ListBlock))
    assert [len(root.items), len(root.items[0].children)] == [2, 1]
    html = PROFILES["clean"].build(document).html
    assert "<ul><li>Parent item<ul><li>Child item</li></ul></li><li>Second parent</li></ul>" in html


def test_inline_formatting(rich_docx: Path):
    html = PROFILES["clean"].build(read_document(rich_docx, rich_docx.name)).html
    assert "<u>underlined</u>" in html
    assert "<s>struck</s>" in html
    assert "H<sub>2</sub>O" in html
    assert "x<sup>2</sup>" in html


def test_hyperlinks_are_kept_only_for_safe_schemes(rich_docx: Path):
    document = read_document(rich_docx, rich_docx.name)
    html = PROFILES["clean"].build(document).html
    assert 'href="https://mindjax.com/lab"' in html
    assert "javascript:" not in html
    assert "hostile link" in html  # label survives without the link
    assert any("unsupported scheme" in w for w in document.warnings)


def test_document_text_is_escaped_not_interpreted(rich_docx: Path):
    html = PROFILES["clean"].build(read_document(rich_docx, rich_docx.name)).html
    assert "<script>" not in html
    assert "&lt;script&gt;alert('x')&lt;/script&gt; &amp; &lt;b&gt;not bold&lt;/b&gt;" in html


def test_tracked_deletions_are_dropped_and_insertions_kept(rich_docx: Path):
    html = PROFILES["clean"].build(read_document(rich_docx, rich_docx.name)).html
    assert "REMOVED-TEXT" not in html
    assert "Kept INSERTED-TEXT" in html


def test_block_content_controls_are_read(rich_docx: Path):
    assert "Inside content control" in PROFILES["clean"].build(read_document(rich_docx, rich_docx.name)).html


def test_merged_cells_use_colspan_and_rowspan(rich_docx: Path):
    document = read_document(rich_docx, rich_docx.name)
    table = next(b for b in document.blocks if isinstance(b, Table))
    assert table.rows[0].cells[0].colspan == 2
    assert table.rows[1].cells[0].rowspan == 2
    assert len(table.rows[2].cells) == 2  # the continued cell is not duplicated
    html = PROFILES["clean"].build(document).html
    assert html.count("Wide header") == 1 and html.count("Tall") == 1
    assert 'colspan="2"' in html and 'rowspan="2"' in html


def test_empty_paragraphs_are_not_emitted(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Before")
    doc.add_paragraph("")
    doc.add_paragraph("   ")
    doc.add_paragraph("After")
    path = tmp_path / "blank.docx"; doc.save(path)
    assert "<p></p>" not in fresh_html(path)


def test_duplicate_images_are_stored_once(tmp_path: Path):
    png = tmp_path / "same.png"
    Image.new("RGB", (400, 300), "red").save(png)
    doc = Document()
    doc.add_paragraph().add_run().add_picture(str(png))
    doc.add_paragraph().add_run().add_picture(str(png))
    path = tmp_path / "dupes.docx"; doc.save(path)
    document = read_document(path, path.name)
    assert len(document.assets) == 1
    assert document.assets[0].occurrences == 2


def test_authored_alt_text_is_used(tmp_path: Path):
    png = tmp_path / "alt.png"
    Image.new("RGB", (400, 300), "blue").save(png)
    doc = Document()
    doc.add_paragraph().add_run().add_picture(str(png))
    doc.inline_shapes[0]._inline.docPr.set("descr", "Login screen with the Continue button highlighted")
    path = tmp_path / "alt.docx"; doc.save(path)
    html = PROFILES["clean"].build(read_document(path, path.name)).html
    assert 'alt="Login screen with the Continue button highlighted"' in html


def test_title_style_becomes_heading(tmp_path: Path):
    doc = Document()
    doc.add_paragraph("Document Title", style="Title")
    path = tmp_path / "title.docx"; doc.save(path)
    document = read_document(path, path.name)
    assert isinstance(document.blocks[0], Heading)
    assert document.title == "Document Title"
