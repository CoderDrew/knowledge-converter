"""Synthetic fixtures. No proprietary documents are needed to run the suite.

`sample_docx` and `sample_pptx` mirror the original FS Knowledge Converter
fixtures so its assertions can be carried over unchanged in spirit.
"""
from __future__ import annotations

import zipfile
from pathlib import Path

import pytest
from docx import Document
from docx.enum.style import WD_STYLE_TYPE
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.util import Inches as PptxInches, Pt

from api.config import Settings


def make_png(path: Path, size=(640, 360), color="#2c5cc5") -> Path:
    image = Image.new("RGB", size, color)
    draw = ImageDraw.Draw(image)
    draw.rectangle((40, 40, size[0] - 40, size[1] - 40), outline="white", width=5)
    draw.text((70, 70), "Synthetic training screenshot", fill="white")
    image.save(path)
    return path


def set_numbering(paragraph, num_id: int, level: int = 0) -> None:
    ppr = paragraph._p.get_or_add_pPr()
    numpr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), str(level))
    numid = OxmlElement("w:numId"); numid.set(qn("w:val"), str(num_id))
    numpr.append(ilvl); numpr.append(numid); ppr.append(numpr)


def add_hyperlink(paragraph, url: str, text: str) -> None:
    rid = paragraph.part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    link = OxmlElement("w:hyperlink"); link.set(qn("r:id"), rid)
    run = OxmlElement("w:r"); t = OxmlElement("w:t"); t.text = text
    run.append(t); link.append(run); paragraph._p.append(link)


@pytest.fixture
def sample_docx(tmp_path: Path) -> Path:
    image = make_png(tmp_path / "shot.png")
    doc = Document()
    doc.add_heading("Cash Management", level=1)
    p = doc.add_paragraph()
    p.add_run("Use ").bold = False
    p.add_run("Cash Management").bold = True
    p.add_run(" to review ").italic = True
    p.add_run("balances.")
    numbered = doc.add_paragraph("Step 1: Open the banking workspace.", style="List Number")
    numbered.add_run().add_picture(str(image), width=Inches(4))
    doc.add_paragraph("Step 2: Select the bank account.", style="List Number")
    doc.add_paragraph("Confirm the displayed balance.", style="List Bullet")
    table = doc.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "Field"; table.cell(0, 1).text = "Value"
    table.cell(1, 0).text = "Status"; table.cell(1, 1).text = "Approved"
    note_style = doc.styles.add_style("Quick Tip", WD_STYLE_TYPE.PARAGRAPH)
    doc.add_paragraph("Quick Tip: verify the date.", style=note_style)
    path = tmp_path / "Cash Management.docx"
    doc.save(path)
    return path


@pytest.fixture
def rich_docx(tmp_path: Path) -> Path:
    """Exercises features the original converter did not handle."""
    doc = Document()
    doc.add_heading("Rich Features", level=1)
    parent = doc.add_paragraph("Parent item", style="List Bullet")
    child = doc.add_paragraph("Child item", style="List Bullet")
    # Word nests list items (Tab key) by raising w:ilvl on the same numbering.
    set_numbering(child, int(parent.style.element.pPr.numPr.numId.val), level=1)
    doc.add_paragraph("Second parent", style="List Bullet")

    p = doc.add_paragraph()
    run = p.add_run("underlined"); run.underline = True
    p.add_run(" ")
    run = p.add_run("struck"); run.font.strike = True
    p.add_run(" H")
    run = p.add_run("2"); run.font.subscript = True
    p.add_run("O and x")
    run = p.add_run("2"); run.font.superscript = True

    links = doc.add_paragraph("Links: ")
    add_hyperlink(links, "https://mindjax.com/lab", "safe link")
    links.add_run(" and ")
    add_hyperlink(links, "javascript:alert(1)", "hostile link")

    doc.add_paragraph("<script>alert('x')</script> & <b>not bold</b>")

    # Tracked change: deleted text must not appear; inserted text must.
    tracked = doc.add_paragraph("Kept ")
    deletion = OxmlElement("w:del"); deletion.set(qn("w:id"), "1"); deletion.set(qn("w:author"), "t")
    del_run = OxmlElement("w:r"); del_text = OxmlElement("w:delText"); del_text.text = "REMOVED-TEXT"
    del_run.append(del_text); deletion.append(del_run); tracked._p.append(deletion)
    insertion = OxmlElement("w:ins"); insertion.set(qn("w:id"), "2"); insertion.set(qn("w:author"), "t")
    ins_run = OxmlElement("w:r"); ins_text = OxmlElement("w:t"); ins_text.text = "INSERTED-TEXT"
    ins_run.append(ins_text); insertion.append(ins_run); tracked._p.append(insertion)

    # Block-level content control (how Word stores e.g. tables of contents).
    sdt = OxmlElement("w:sdt"); content = OxmlElement("w:sdtContent")
    sdt_p = OxmlElement("w:p"); sdt_r = OxmlElement("w:r"); sdt_t = OxmlElement("w:t"); sdt_t.text = "Inside content control"
    sdt_r.append(sdt_t); sdt_p.append(sdt_r); content.append(sdt_p); sdt.append(content)
    doc.element.body.insert(len(doc.element.body) - 1, sdt)

    table = doc.add_table(rows=3, cols=3)
    merged = table.cell(0, 0).merge(table.cell(0, 1))
    merged.text = "Wide header"
    table.cell(0, 2).text = "Right"
    tall = table.cell(1, 0).merge(table.cell(2, 0))
    tall.text = "Tall"
    table.cell(1, 1).text = "b1"; table.cell(1, 2).text = "c1"
    table.cell(2, 1).text = "b2"; table.cell(2, 2).text = "c2"

    path = tmp_path / "Rich Features.docx"
    doc.save(path)
    return path


@pytest.fixture
def sample_pptx(tmp_path: Path) -> Path:
    image = make_png(tmp_path / "slide.png")
    prs = Presentation()
    slide = prs.slides.add_slide(prs.slide_layouts[5])
    slide.shapes.title.text = "Cash Management"
    slide.shapes.add_picture(str(image), PptxInches(1), PptxInches(1.5), width=PptxInches(5))
    path = tmp_path / "STP - Cash Management.pptx"
    prs.save(path)
    return path


@pytest.fixture
def rich_pptx(tmp_path: Path) -> Path:
    image = make_png(tmp_path / "deck.png")
    prs = Presentation()
    title = prs.slides.add_slide(prs.slide_layouts[0])
    title.shapes.title.text = "Quarterly Review"
    title.placeholders[1].text = "Prepared by the operations team"

    bullets = prs.slides.add_slide(prs.slide_layouts[1])
    bullets.shapes.title.text = "Agenda"
    frame = bullets.placeholders[1].text_frame
    frame.text = "Results"
    sub = frame.add_paragraph(); sub.text = "Revenue detail"; sub.level = 1
    nxt = frame.add_paragraph(); nxt.text = "Next steps"
    bullets.notes_slide.notes_text_frame.text = "Remember to pause for questions."

    # A deck built only from text boxes (no title placeholder), like many real decks.
    free = prs.slides.add_slide(prs.slide_layouts[6])
    lower = free.shapes.add_textbox(PptxInches(1), PptxInches(3), PptxInches(6), PptxInches(1))
    lower.text_frame.text = "Body copy placed lower on the slide."
    heading = free.shapes.add_textbox(PptxInches(1), PptxInches(0.4), PptxInches(8), PptxInches(1))
    heading.text_frame.text = "Free-form Title"
    heading.text_frame.paragraphs[0].runs[0].font.size = Pt(36)
    free.shapes.add_picture(str(image), PptxInches(1), PptxInches(4.2), width=PptxInches(3))
    group = free.shapes.add_group_shape()
    grouped = group.shapes.add_textbox(PptxInches(6), PptxInches(5), PptxInches(3), PptxInches(1))
    grouped.text_frame.text = "Grouped caption"

    data = prs.slides.add_slide(prs.slide_layouts[5])
    data.shapes.title.text = "Figures"
    shape = data.shapes.add_table(3, 2, PptxInches(1), PptxInches(1.5), PptxInches(6), PptxInches(2))
    tbl = shape.table
    tbl.cell(0, 0).text = "Metric"; tbl.cell(0, 1).text = "Value"
    tbl.cell(1, 0).merge(tbl.cell(1, 1)); tbl.cell(1, 0).text = "Merged row"
    tbl.cell(2, 0).text = "Uptime"; tbl.cell(2, 1).text = "99.9%"

    path = tmp_path / "Quarterly Review.pptx"
    prs.save(path)
    return path


def rezip(source: Path, destination: Path, *, replace: dict[str, bytes] | None = None, extra: dict[str, bytes] | None = None) -> Path:
    """Copy an OOXML package, replacing or adding members."""
    replace = replace or {}
    with zipfile.ZipFile(source) as src, zipfile.ZipFile(destination, "w", zipfile.ZIP_DEFLATED) as dst:
        for info in src.infolist():
            dst.writestr(info.filename, replace.get(info.filename, src.read(info.filename)))
        for name, data in (extra or {}).items():
            dst.writestr(name, data)
    return destination


@pytest.fixture
def settings(tmp_path: Path) -> Settings:
    return Settings(work_dir=tmp_path / "work", isolation="inline", rate_limit_per_minute=0)
